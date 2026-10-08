#!/usr/bin/env python3
"""Genera los archivos de descarga del Archivo abierto, ya maquetados:

  descargas/fichas/Ficha_NNN.pdf     cada ficha completa, con sus fotografías
  descargas/indice_fichas.pdf        índice de fichas
  descargas/personas.pdf             índice de personas
  descargas/biblioteca.pdf           biblioteca
  descargas/atlas_giles_datos.xlsx   todos los datos en un libro de Excel

Uso (desde la raíz del repositorio, después de construir_archivo.py):
  python3 tools/generar_descargas.py            todo
  python3 tools/generar_descargas.py 048 049    sólo esas fichas (más índices y Excel)
Requiere openpyxl, Pillow y un Chromium o Chrome (variable CHROMIUM, o los que encuentre en el sistema).
"""
import sys, os, re, json, glob, html, shutil, subprocess, tempfile, datetime
from PIL import Image
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
A = os.path.join(ROOT, 'data', 'archivo')
OUT = os.path.join(ROOT, 'descargas')
FU = os.path.join(ROOT, 'tools', 'fuentes')
HOY = datetime.date.today()
MESES = 'enero febrero marzo abril mayo junio julio agosto septiembre octubre noviembre diciembre'.split()
FECHA = '%d de %s de %d' % (HOY.day, MESES[HOY.month - 1], HOY.year)
SITIO = 'https://surtecto.github.io/atlas-mensuras-giles/'
# paleta de las fichas en Word (tema «Rojo Naranja») y tipografía Georgia / Gelasio
C = dict(osc='B22600', ac='E84C22', oro='FFBD47', tie='B64926', fon='FBE3D6', tx='262626', gr='505046')

def J(p, d=None):
    try:
        with open(p, encoding='utf8') as f: return json.load(f)
    except FileNotFoundError: return d
e = lambda s: html.escape(str(s if s is not None else ''))
def pal(s):
    s = e(s)
    s = re.sub(r'\^\{([^}]*)\}', r'<sup>\1</sup>', s); s = re.sub(r'~~(.*?)~~', r'<del>\1</del>', s)
    s = re.sub(r'\*\*(.*?)\*\*', r'<b>\1</b>', s); return re.sub(r'\[([^\]]*)\]', r'<i class="ed">[\1]</i>', s)
plano = lambda s: re.sub(r'\*\*', '', re.sub(r'~~(.*?)~~', r'⟨\1⟩', re.sub(r'\^\{([^}]*)\}', r'\1', s or '')))
par = lambda s: ''.join('<p>%s</p>' % e(x) for x in str(s or '').split('\n') if x.strip())

def chromium():
    c = [os.environ.get('CHROMIUM'), '/opt/pw-browsers/chromium'] + sorted(glob.glob('/opt/pw-browsers/chromium-*/chrome-linux/chrome')) + \
        [shutil.which(x) for x in ('chromium', 'chromium-browser', 'google-chrome', 'chrome', 'msedge')] + \
        [r'C:\Program Files\Google\Chrome\Application\chrome.exe', r'C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe',
         '/Applications/Google Chrome.app/Contents/MacOS/Google Chrome']
    for x in c:
        if x and os.path.exists(x): return x
    raise SystemExit('No encuentro Chromium/Chrome. Indicá su ruta en la variable CHROMIUM.')

CSS = '''
@font-face{font-family:G;src:url("file://%(fu)s/gelasio-latin-400-normal.woff2") format("woff2");font-weight:400}
@font-face{font-family:G;src:url("file://%(fu)s/gelasio-latin-400-italic.woff2") format("woff2");font-weight:400;font-style:italic}
@font-face{font-family:G;src:url("file://%(fu)s/gelasio-latin-700-normal.woff2") format("woff2");font-weight:700}
@font-face{font-family:G;src:url("file://%(fu)s/gelasio-latin-700-italic.woff2") format("woff2");font-weight:700;font-style:italic}
@page{size:A4 %(ori)s;margin:17mm 15mm 17mm 15mm;
  @bottom-left{content:"%(pie)s";font:7.5pt G,Georgia,serif;color:#505046}
  @bottom-right{content:counter(page) " / " counter(pages);font:7.5pt G,Georgia,serif;color:#505046}}
@page:first{margin-top:14mm}
*{box-sizing:border-box}
html{font:10pt/1.42 G,Georgia,serif;color:#262626}
body{margin:0}
p{margin:0 0 2.2mm}
h1,h2,h3,h4{margin:0;font-weight:700}
.sup{font-size:7.5pt;letter-spacing:.14em;text-transform:uppercase;color:#B64926;font-weight:700}
.tit{font-size:25pt;line-height:1.08;color:#B22600;margin:1.5mm 0 1.5mm}
.sub{font-size:10.5pt;color:#505046;font-style:italic;margin-bottom:5mm}
.banda{background:#B22600;color:#fff;padding:2.6mm 4mm;display:flex;gap:4mm;align-items:baseline;margin:0 0 4mm;break-after:avoid}
.banda .n{font-size:15pt;font-weight:700;color:#FFBD47;font-variant-numeric:tabular-nums}
.banda h2{font-size:11.5pt;letter-spacing:.03em;flex:1}
.banda span{font-size:7.5pt;color:#FFE6C7;text-align:right}
.sec{break-before:page}
h3{font-size:8pt;letter-spacing:.12em;text-transform:uppercase;color:#E84C22;border-bottom:.6pt solid #E84C22;padding-bottom:.8mm;margin:4mm 0 2mm;break-after:avoid}
table{border-collapse:collapse;width:100%%}
table.kv td{padding:1.4mm 2.5mm;vertical-align:top;border-bottom:.4pt solid #E9C9B7;font-size:9.3pt}
table.kv td:first-child{width:38mm;background:#FBE3D6;color:#B22600;font-weight:700;font-size:8.2pt}
.dos{display:flex;gap:6mm;align-items:flex-start}
.dos>*{min-width:0}
figure{margin:0;break-inside:avoid}
figure img{display:block;max-width:100%%;border:.5pt solid #B64926}
figcaption{font-size:7.3pt;color:#505046;font-style:italic;margin-top:1mm}
.caja{background:#FBE3D6;border-left:2.2pt solid #E84C22;padding:3mm 4mm;margin:4mm 0;break-inside:avoid}
.caja p:last-child{margin:0}
ol.tx{list-style:none;margin:0;padding:0;counter-reset:r;font-size:9.4pt;line-height:1.36}
ol.tx li{counter-increment:r;display:flex;gap:2.2mm;break-inside:avoid}
ol.tx li::before{content:counter(r);flex:0 0 6mm;text-align:right;font-size:6.6pt;line-height:2;color:#B64926;font-variant-numeric:tabular-nums}
ol.tx sup{font-size:.68em;line-height:0}
ol.tx del{color:#8a8a80}
.ed{color:#6b6b60}
.nota{font-size:8.3pt;color:#505046;font-style:italic;margin-bottom:2.5mm}
.interp{margin-top:4mm}
.interp p{text-align:justify;hyphens:auto}
.obs h4,.bio h4{font-size:10pt;color:#B22600;margin:3.2mm 0 1mm;break-after:avoid}
.bio h4 small{font-size:7pt;letter-spacing:.1em;text-transform:uppercase;color:#E84C22;margin-left:2mm}
.bio .f{font-size:7.8pt;color:#505046;font-style:italic}
.bio>div,.obs>div{break-inside:avoid}
.cols{columns:2;column-gap:7mm}
.cita{font-size:8.3pt;color:#505046;border-top:.6pt solid #B64926;padding-top:2.5mm;margin-top:6mm;break-inside:avoid}
table.lst{font-size:8.3pt}
table.lst th{background:#B22600;color:#fff;text-align:left;padding:1.8mm 2mm;font-size:7.3pt;letter-spacing:.08em;text-transform:uppercase}
table.lst td{padding:1.8mm 2mm;vertical-align:top;border-bottom:.4pt solid #E9C9B7}
table.lst tr:nth-child(even) td{background:#FDF3EC}
table.lst tr{break-inside:avoid}
table.lst thead{display:table-header-group}
td.n{font-weight:700;color:#B22600;white-space:nowrap;font-variant-numeric:tabular-nums}
.cifras{display:flex;gap:9mm;margin:4mm 0 5mm}
.cifras div{font-size:8pt;color:#505046}.cifras b{display:block;font-size:17pt;color:#B22600;line-height:1.1}
'''

def doc(titulo, cuerpo, pie, ori='portrait'):
    return '<!doctype html><html lang="es"><head><meta charset="utf-8"><title>%s</title><style>%s</style></head><body>%s</body></html>' % (
        e(titulo), CSS % dict(fu=FU.replace('\\', '/'), ori=ori, pie=pie.replace('"', '”')), cuerpo)

def a_pdf(html_txt, destino, tmp):
    src = os.path.join(tmp, 'p.html')
    with open(src, 'w', encoding='utf8') as f: f.write(html_txt)
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    subprocess.run([chromium(), '--headless=new', '--no-sandbox', '--disable-gpu', '--no-pdf-header-footer', '--allow-file-access-from-files',
                    '--virtual-time-budget=8000', '--print-to-pdf=' + destino, 'file://' + src], check=True, capture_output=True, timeout=300)

def foto(rel, tmp, ancho=900, q=58):
    """copia reducida para el PDF (las de la web quedan a mayor tamaño)"""
    if not rel: return ''
    src = os.path.join(ROOT, rel)
    if not os.path.exists(src): return ''
    dst = os.path.join(tmp, rel.replace('/', '_'))
    if not os.path.exists(dst):
        im = Image.open(src); im.thumbnail((ancho, ancho), Image.LANCZOS); im.save(dst, 'JPEG', quality=q, optimize=True)
    return 'file://' + dst

def fig(rel, pie, tmp, alto='118mm'):
    u = foto(rel, tmp)
    return '<figure><img src="%s" style="max-height:%s"><figcaption>%s</figcaption></figure>' % (u, alto, e(pie)) if u else ''

def cita(d, n): return ('Addesso, Juan Patricio (%d). «Ficha %03d: %s (%s)». Transcripción e interpretación del duplicado de mensura n.º %d del partido de '
    'San Andrés de Giles, Archivo Histórico de Geodesia de la Provincia de Buenos Aires. Atlas de Mensuras de Giles, Archivo abierto. %sarchivo.html#f%03d') % (
    HOY.year, n, d.get('titulo', ''), d.get('resumen', {}).get('Año', ''), n, SITIO, n)

def ficha_html(d, tmp):
    n = int(d['carpeta']); R = d.get('resumen', {}); M = d.get('metricas', {}); k = [0]
    def banda(t, s=''):
        k[0] += 1; return '<div class="banda"><span class="n">%02d</span><h2>%s</h2><span>%s</span></div>' % (k[0], e(t), e(s))
    kv = lambda filas: '<table class="kv">%s</table>' % ''.join('<tr><td>%s</td><td>%s</td></tr>' % (e(a), e(b)) for a, b in filas if b)
    car = d.get('caratula', {})
    h = ['<div class="sup">Archivo Histórico de Geodesia · Provincia de Buenos Aires</div><div class="sup" style="color:#E84C22;margin-top:1mm">Ficha de mensura n.º %03d</div>' % n,
         '<h1 class="tit">%s</h1><div class="sub">%s · %s · Agrim. %s</div>' % (e(d.get('titulo', '')), e(R.get('Ubicación', '')), e(R.get('Año', '')), e(R.get('Agrimensor', ''))),
         '<div class="dos"><div style="flex:1.55">%s</div><div style="flex:1">%s</div></div>' % (kv(R.items()), fig(car.get('img'), 'Carátula · ' + car.get('fuente', ''), tmp, '112mm')),
         '<div class="caja"><div class="sup" style="margin-bottom:1.5mm">Síntesis interpretativa</div>%s</div>' % par(d.get('sintesis', '')),
         kv([('Superficie (ha)', M.get('superficie_ha')), ('Hitos del paisaje', M.get('hitos')), ('Cursos de agua', M.get('aguas'))]),
         '<p class="nota" style="margin-top:4mm">Transcripción paleográfica: se respetan grafía, puntuación y renglones del original, numerados. Letras voladas en superíndice; '
         '<del>testado</del>; <i class="ed">[entre corchetes]</i>, nota o reconstrucción del transcriptor; <i class="ed">[?]</i>, lectura dudosa.</p>']
    if car.get('campos'):
        h.append('<section class="sec">%s<div class="dos"><div style="flex:1.4">%s</div><div style="flex:1">%s</div></div>' % (
            banda('Carátula', car.get('fuente', '')), kv(car['campos']), fig(car.get('img'), 'Facsímil de la carátula', tmp, '150mm')))
        po = d.get('portadilla')
        if po and po.get('lineas'):
            h.append('<h3>Portadilla manuscrita</h3><div class="dos"><div style="flex:1.3"><ol class="tx">%s</ol></div><div style="flex:1">%s</div></div>' % (
                ''.join('<li><span>%s</span></li>' % pal(l) for l in po['lineas']), fig(po.get('img'), 'Portadilla · ' + po.get('fuente', ''), tmp, '95mm')))
        h.append('</section>')
    for f in d.get('folios', []):
        h.append('<section class="sec">%s%s<div class="dos"><div style="flex:1.45"><ol class="tx">%s</ol></div><div style="flex:1">%s</div></div>'
                 '<div class="interp"><h3>Interpretación · ¿de qué se habla en esta hoja?</h3>%s</div></section>' % (
            banda(f.get('titulo', ''), f.get('fuente', '')), '<p class="nota">%s</p>' % e(f['nota']) if f.get('nota') else '',
            ''.join('<li><span>%s</span></li>' % pal(l) for l in f.get('lineas', [])), fig(f.get('img'), 'Facsímil · ' + f.get('fuente', ''), tmp, '150mm'), par(f.get('interpretacion', ''))))
    for i, p in enumerate(d.get('planos', [])):
        h.append('<section class="sec">%s%s<h3>Lectura del plano</h3><div class="interp" style="margin-top:0">%s</div>%s</section>' % (
            banda('Planimetría · lámina %d' % (i + 1), p.get('titulo', '')), fig(p.get('img'), 'Fuente: %s%s' % (p.get('fuente', ''), ' — ' + p['orientacion'] if p.get('orientacion') else ''), tmp, '165mm'),
            par(p.get('interpretacion', '')), '<h3>Referencias</h3>' + kv(p['referencias']) if p.get('referencias') else ''))
    if d.get('observaciones'):
        h.append('<section class="sec">%s<div class="obs">%s</div></section>' % (banda('Observaciones de lectura', 'paisaje · ritual · técnica · tenencia'),
            ''.join('<div><h4>%s</h4>%s</div>' % (e(o['titulo']), par(o['texto'])) for o in d['observaciones'])))
    if d.get('biografias'):
        h.append('<section class="sec">%s<div class="bio cols">%s</div></section>' % (banda('Personas nombradas', 'notas biográficas'),
            ''.join('<div><h4>%s<small>%s</small></h4>%s<p class="f">Fuentes: %s</p></div>' % (e(b['nombre']), e(b['rol']), par(b['texto']), e(b['fuentes'])) for b in d['biografias'])))
    h.append('<div class="cita"><b>Cómo citar.</b> %s (generado el %s).<br>Transcripción y lectura bajo licencia CC BY 4.0. Las reproducciones de los documentos pertenecen al Archivo Histórico de Geodesia.</div>' % (e(cita(d, n)), FECHA))
    return doc('Ficha %03d · %s' % (n, d.get('titulo', '')), ''.join(h), 'Atlas de Mensuras de Giles · Ficha %03d · %s' % (n, d.get('titulo', '')))

def cabecera(sup, tit, sub): return '<div class="sup">%s</div><h1 class="tit">%s</h1><div class="sub">%s</div>' % (e(sup), e(tit), e(sub))

def indice_html(F, tot):
    ren = sum(f.get('renglones', 0) for f in F); fol = sum(f.get('folios', 0) for f in F)
    c = cabecera('Atlas de Mensuras de Giles · Archivo abierto', 'Índice de fichas de mensura',
                 'Duplicados de mensura del partido de San Andrés de Giles · Archivo Histórico de Geodesia de la Provincia de Buenos Aires · actualizado el ' + FECHA)
    c += '<div class="cifras"><div><b>%d</b>fichas de %d carpetas</div><div><b>%d</b>hojas transcriptas</div><div><b>%s</b>renglones</div><div><b>%s–%s</b>años de las mensuras</div></div>' % (
        len(F), tot, fol, '{:,}'.format(ren).replace(',', '.'), min(f['y'][:4] for f in F if f.get('y')), max(f['y'][:4] for f in F if f.get('y')))
    c += '<table class="lst" style="table-layout:fixed;font-size:7.9pt;line-height:1.34"><colgroup><col style="width:10mm"><col style="width:35mm"><col style="width:13mm"><col style="width:25mm"><col style="width:36mm"><col style="width:23mm"><col style="width:11mm"><col></colgroup><thead><tr><th>N.º</th><th>Titular</th><th>Año</th><th>Agrimensor</th><th>Paraje</th><th>Sup. (ha)</th><th>Hojas</th><th>Síntesis</th></tr></thead><tbody>'
    for f in F:
        c += '<tr><td class="n">%03d</td><td><b>%s</b></td><td>%s</td><td>%s</td><td>%s</td><td>%s</td><td>%s</td><td>%s</td></tr>' % (
            f['n'], e(f.get('prop')), e(f.get('y')), e(f.get('ag')), e(f.get('ubic')), e(f.get('sup')), f.get('folios', ''), e(f.get('sint')))
    return doc('Índice de fichas', c + '</tbody></table>', 'Atlas de Mensuras de Giles · Índice de fichas · ' + SITIO + 'archivo.html', 'landscape')

def personas_html(P, byn):
    c = cabecera('Atlas de Mensuras de Giles · Archivo abierto', 'Personas nombradas en las mensuras', '%d notas biográficas de propietarios, linderos, agrimensores y autoridades · actualizado el %s' % (len(P), FECHA))
    c += '<div class="bio cols">' + ''.join('<div><h4>%s<small>%s</small></h4>%s<p class="f">Ficha %03d · %s. Fuentes: %s</p></div>' % (
        e(p['nombre']), e(p['rol']), par(p['texto']), p['n'], e(byn.get(p['n'], {}).get('prop', '')), e(p['fuentes'])) for p in P) + '</div>'
    return doc('Personas', c, 'Atlas de Mensuras de Giles · Personas · ' + SITIO + 'archivo.html#personas')

TB = dict(libro='Libro', articulo='Artículo', tesis='Tesis', ponencia='Ponencia', capitulo='Capítulo', fuente='Fuente primaria', recurso='Recurso digital', informe='Informe', mapa='Cartografía')
def biblio_html(B):
    c = cabecera('Atlas de Mensuras de Giles · Archivo abierto', 'Biblioteca del territorio', 'Investigaciones, libros y fuentes sobre San Andrés de Giles y su región · %d referencias · actualizado el %s' % (len(B), FECHA))
    c += '<table class="lst"><thead><tr><th>Año</th><th style="width:38mm">Autor</th><th>Título y publicación</th><th>Tipo</th><th style="width:48mm">Enlace</th></tr></thead><tbody>' + ''.join(
        '<tr><td class="n">%s</td><td>%s</td><td><b>%s</b><br>%s%s</td><td>%s</td><td style="word-break:break-all;font-size:7pt">%s</td></tr>' % (
            e(b['anio']), e('; '.join(b.get('autores', []))), e(b['titulo']), e(b.get('en', '')), '<br><i>%s</i>' % e(b['nota']) if b.get('nota') else '', e(TB.get(b['tipo'], b['tipo'])), e(b.get('url', ''))) for b in B) + '</tbody></table>'
    return doc('Biblioteca', c, 'Atlas de Mensuras de Giles · Biblioteca · ' + SITIO + 'archivo.html#biblioteca')

# ---------------------------------------------------------------- Excel
def hoja(wb, nombre, cols, filas, alto=None):
    """cols = [(título, ancho, alineación)]"""
    ws = wb.create_sheet(nombre)
    fino = Side(style='thin', color='E9C9B7')
    for j, (t, w, *_) in enumerate(cols, 1):
        c = ws.cell(1, j, t); c.font = Font(name='Georgia', size=10, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=C['osc'])
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = Border(bottom=Side(style='medium', color=C['ac']))
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 30
    for i, fila in enumerate(filas, 2):
        for j, v in enumerate(fila, 1):
            c = ws.cell(i, j, v); al = cols[j - 1][2] if len(cols[j - 1]) > 2 else 'left'
            c.font = Font(name='Georgia', size=10, bold=(j == 1), color=C['osc'] if j == 1 else C['tx'])
            c.alignment = Alignment(horizontal=al, vertical='top', wrap_text=True); c.border = Border(bottom=fino)
            if i % 2: c.fill = PatternFill('solid', fgColor='FDF3EC')
        if alto: ws.row_dimensions[i].height = alto
    ws.freeze_panes = 'C2' if len(cols) > 3 else 'A2'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(cols)), max(len(filas) + 1, 2))
    ws.sheet_view.zoomScale = 100; ws.sheet_properties.tabColor = C['ac']
    ws.page_setup.orientation = 'landscape'; ws.page_setup.paperSize = 9; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True; ws.print_title_rows = '1:1'
    ws.oddFooter.center.text = 'Atlas de Mensuras de Giles · ' + nombre + ' · pág. &P de &N'
    return ws

def excel(F, D, P, B, AP, destino):
    wb = Workbook(); L = wb.active; L.title = 'Léame'; L.sheet_properties.tabColor = C['osc']; L.sheet_view.showGridLines = False
    L.column_dimensions['A'].width = 3; L.column_dimensions['B'].width = 24; L.column_dimensions['C'].width = 96
    L['B2'] = 'ATLAS DE MENSURAS DE GILES · ARCHIVO ABIERTO'; L['B2'].font = Font(name='Georgia', size=9, bold=True, color=C['tie'])
    L['B3'] = 'Datos de las fichas de mensura'; L['B3'].font = Font(name='Georgia', size=22, bold=True, color=C['osc']); L.row_dimensions[3].height = 34
    L['B4'] = 'Duplicados de mensura del partido de San Andrés de Giles · Archivo Histórico de Geodesia de la Provincia de Buenos Aires'; L['B4'].font = Font(name='Georgia', size=10, italic=True, color=C['gr'])
    info = [('Hoja', 'Contenido'), ('Fichas', '%d fichas: una fila por carpeta, con titular, agrimensor, año, paraje, medidas, linderos, hitos, aguas y síntesis.' % len(F)),
            ('Hojas', 'Una fila por hoja transcripta: título, fuente fotográfica, cantidad de renglones e interpretación.'),
            ('Renglones', 'La transcripción paleográfica completa, un renglón por fila (ficha, hoja, número de renglón, texto). Sirve para buscar y filtrar.'),
            ('Láminas', 'Una fila por plano o detalle: título, fuente, orientación y lectura.'), ('Personas', '%d notas biográficas de las personas nombradas.' % len(P)),
            ('Biblioteca', '%d referencias sobre el territorio.' % len(B)), ('Aportes', 'Aportes del público ya publicados.'), ('', ''),
            ('Notación', 'En «Renglones»: ⟨texto⟩ = testado en el original; [texto] = nota o reconstrucción del transcriptor; [?] = lectura dudosa. Se respeta la ortografía del original.'),
            ('Actualizado', FECHA), ('Autor', 'Juan Patricio Addesso · Doctorado en Geografía, Universidad del Salvador · Surtectura'),
            ('Cómo citar', 'Addesso, Juan Patricio (%d). Atlas de Mensuras de Giles: archivo abierto. %s' % (HOY.year, SITIO)),
            ('Licencia', 'Transcripciones, lecturas y datos: CC BY 4.0. Los documentos originales pertenecen al Archivo Histórico de Geodesia.')]
    for i, (a, b) in enumerate(info, 6):
        ca, cb = L.cell(i, 2, a), L.cell(i, 3, b); cb.alignment = Alignment(wrap_text=True, vertical='top'); ca.alignment = Alignment(vertical='top')
        ca.font = Font(name='Georgia', size=10, bold=True, color=C['osc']); cb.font = Font(name='Georgia', size=10, color=C['tx'])
        if i == 6:
            for c in (ca, cb): c.font = Font(name='Georgia', size=10, bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=C['osc'])
        elif a and i < 14: ca.fill = PatternFill('solid', fgColor=C['fon'])
        if len(b) > 95: L.row_dimensions[i].height = 30
    hoja(wb, 'Fichas', [('Ficha', 7, 'center'), ('Titular', 30), ('Año', 7, 'center'), ('Fecha de la diligencia', 22), ('Agrimensor', 22), ('Paraje', 34), ('Superficie (ha)', 18),
                        ('Medidas', 34), ('Rumbos y declinación', 30), ('Orden judicial', 30), ('Vendedor o antecesor', 26), ('Linderos', 38), ('Hitos del paisaje', 38), ('Cursos de agua', 30),
                        ('Referencia de archivo', 30), ('Hojas', 7, 'center'), ('Renglones', 10, 'center'), ('Láminas', 8, 'center'), ('Fotos', 7, 'center'), ('Síntesis', 90), ('PDF', 44)],
         [['%03d' % f['n'], f.get('prop'), f.get('y'), f.get('fecha'), f.get('ag'), f.get('ubic'), f.get('sup'), D[f['n']]['resumen'].get('Medidas') if f['n'] in D else '',
           D[f['n']]['resumen'].get('Rumbos') if f['n'] in D else '', D[f['n']]['resumen'].get('Orden judicial') if f['n'] in D else '', D[f['n']]['resumen'].get('Vendedor / antecesor') if f['n'] in D else '',
           f.get('lind'), f.get('hitos'), f.get('agua'), f.get('ref'), f.get('folios'), f.get('renglones'), f.get('planos'), f.get('fotos'), f.get('sint'),
           '%sdescargas/fichas/Ficha_%03d.pdf' % (SITIO, f['n'])] for f in F], alto=95)
    hoja(wb, 'Hojas', [('Ficha', 7, 'center'), ('Titular', 26), ('N.º', 5, 'center'), ('Hoja', 44), ('Fuente', 24), ('Renglones', 10, 'center'), ('Nota', 40), ('Interpretación', 110)],
         [['%03d' % n, d.get('titulo'), i, f.get('titulo'), f.get('fuente'), len(f.get('lineas', [])), f.get('nota'), f.get('interpretacion')] for n, d in sorted(D.items()) for i, f in enumerate(d.get('folios', []), 1)], alto=80)
    hoja(wb, 'Renglones', [('Ficha', 7, 'center'), ('Año', 7, 'center'), ('Hoja', 40), ('Renglón', 9, 'center'), ('Texto', 95)],
         [['%03d' % n, d['resumen'].get('Año'), f.get('titulo'), i, plano(l)] for n, d in sorted(D.items()) for f in d.get('folios', []) for i, l in enumerate(f.get('lineas', []), 1)])
    hoja(wb, 'Láminas', [('Ficha', 7, 'center'), ('Titular', 26), ('Lámina', 8, 'center'), ('Título', 44), ('Fuente', 20), ('Orientación y soporte', 44), ('Lectura del plano', 110)],
         [['%03d' % n, d.get('titulo'), i, p.get('titulo'), p.get('fuente'), p.get('orientacion'), p.get('interpretacion')] for n, d in sorted(D.items()) for i, p in enumerate(d.get('planos', []), 1)], alto=80)
    hoja(wb, 'Personas', [('Nombre', 34), ('Rol', 26), ('Ficha', 7, 'center'), ('Nota biográfica', 110), ('Fuentes', 40)], [[p['nombre'], p['rol'], '%03d' % p['n'], p['texto'], p['fuentes']] for p in P], alto=66)
    hoja(wb, 'Biblioteca', [('Año', 9, 'center'), ('Autores', 30), ('Título', 60), ('Publicado en', 46), ('Tipo', 16), ('Acceso', 16), ('Temas', 30), ('Enlace', 50), ('Nota', 50)],
         [[b['anio'], '; '.join(b.get('autores', [])), b['titulo'], b.get('en'), TB.get(b['tipo'], b['tipo']), b.get('acceso'), '; '.join(b.get('temas', [])), b.get('url'), b.get('nota')] for b in B], alto=45)
    hoja(wb, 'Aportes', [('Fecha', 12, 'center'), ('Tipo', 16), ('Título', 40), ('Descripción', 70), ('Lugar', 26), ('Coordenadas', 20), ('Época', 12), ('Carpeta', 9, 'center'), ('Aportado por', 24), ('Fuente', 28), ('Enlace', 40)],
         [[a.get('fecha'), a.get('tipo'), a.get('titulo'), a.get('descripcion'), a.get('lugar'), a.get('coordenadas'), a.get('anio'), a.get('carpeta'), a.get('autor'), a.get('fuente'), a.get('url')] for a in AP], alto=45)
    wb.save(destino)

if __name__ == '__main__':
    cuales = set(int(x) for x in sys.argv[1:])
    IDX = J(os.path.join(A, 'fichas.json')); F = IDX['fichas']; byn = {f['n']: f for f in F}
    D = {int(d['carpeta']): d for d in (J(p) for p in sorted(glob.glob(os.path.join(A, 'fichas', '*.json'))))}
    P, B, AP = J(os.path.join(A, 'personas.json'), []), J(os.path.join(A, 'biblioteca.json'), []), J(os.path.join(A, 'aportes.json'), [])
    os.makedirs(OUT, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        for n, d in sorted(D.items()):
            if cuales and n not in cuales: continue
            dst = os.path.join(OUT, 'fichas', 'Ficha_%03d.pdf' % n)
            a_pdf(ficha_html(d, tmp), dst, tmp); print('Ficha %03d  %5.1f MB' % (n, os.path.getsize(dst) / 1048576), flush=True)
        a_pdf(indice_html(F, IDX.get('total_corpus', 191)), os.path.join(OUT, 'indice_fichas.pdf'), tmp)
        a_pdf(personas_html(P, byn), os.path.join(OUT, 'personas.pdf'), tmp)
        a_pdf(biblio_html(B), os.path.join(OUT, 'biblioteca.pdf'), tmp)
    excel(F, D, P, B, AP, os.path.join(OUT, 'atlas_giles_datos.xlsx'))
    print('Índices en PDF y libro de Excel listos. Volvé a correr construir_archivo.py para registrar los PDF nuevos.')
